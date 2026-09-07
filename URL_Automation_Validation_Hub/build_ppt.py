#!/usr/bin/env python3
"""Build a 3-slide Synchrony-branded PPT: title, approach, future."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
LOGO_PNG = ASSETS / "synchrony_logo.png"
OUTPUT = ROOT / "URL_Validation_Approach.pptx"

GOLD = RGBColor(0xF9, 0xC2, 0x0A)
CHARCOAL = RGBColor(0x34, 0x37, 0x41)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE = RGBColor(0xF6, 0xF5, 0xF1)
MUTED = RGBColor(0x6B, 0x6E, 0x76)
CARD_BORDER = RGBColor(0xEE, 0xEE, 0xEA)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def make_logo() -> None:
    """Draw the Synchrony lockup: three gold pillars + wordmark."""
    ASSETS.mkdir(exist_ok=True)
    scale = 8
    width, height = 300 * scale, 70 * scale
    img = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    gold = (249, 194, 10, 255)
    # SVG proportions, scaled
    draw.rectangle([1 * scale, 32.3 * scale, 13.3 * scale, 64.1 * scale], fill=gold)
    draw.rectangle([19.1 * scale, 0.5 * scale, 31.4 * scale, 64.1 * scale], fill=gold)
    draw.rectangle([37.2 * scale, 0.5 * scale, 49.4 * scale, 32.3 * scale], fill=gold)

    font_path = None
    for candidate in (
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ):
        if Path(candidate).exists():
            font_path = candidate
            break
    font = ImageFont.truetype(font_path, 34 * scale) if font_path else ImageFont.load_default()
    draw.text((62 * scale, 18 * scale), "synchrony", font=font, fill=(52, 55, 65, 255))
    # Flatten onto white so the logo is readable on every slide
    bg = Image.new("RGBA", img.size, (255, 255, 255, 255))
    composed = Image.alpha_composite(bg, img).convert("RGB")
    composed.save(LOGO_PNG)


def set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"


def add_textbox(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(6)
        run = p.runs[0] if p.runs else p.add_run()
        set_run(run, line, size, color, bold)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, l, t, w, h, color, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, l, t, w, h)
    fill_shape(shape, color)
    if rounded:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def add_card(slide, l, t, w, h):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(0.75)
    try:
        shape.adjustments[0] = 0.05
    except Exception:
        pass
    return shape


def add_logo(slide, left, top, width):
    with Image.open(LOGO_PNG) as img:
        ratio = img.height / img.width
    height = Emu(int(width * ratio))
    slide.shapes.add_picture(str(LOGO_PNG), left, top, width=width, height=height)
    return height


def add_footer(slide, page, total=3):
    add_rect(slide, Inches(0), Inches(7.28), SLIDE_W, Inches(0.22), CHARCOAL)
    add_rect(slide, Inches(0), Inches(7.28), Inches(0.42), Inches(0.22), GOLD)
    add_textbox(
        slide, Inches(0.55), Inches(7.28), Inches(8.5), Inches(0.22),
        "Synchrony  |  URL Validation", 10, WHITE, False, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide, Inches(11.4), Inches(7.28), Inches(1.6), Inches(0.22),
        f"{page}  /  {total}", 10, WHITE, False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE,
    )


def set_blank_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def build_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_blank_bg(slide, WHITE)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.12), GOLD)
    add_rect(slide, Inches(0), Inches(7.38), SLIDE_W, Inches(0.12), GOLD)

    add_logo(slide, Inches(5.05), Inches(2.05), Inches(3.2))
    add_textbox(
        slide, Inches(0.8), Inches(3.45), Inches(11.7), Inches(1.2),
        "URL Validation", 48, CHARCOAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, Inches(6.0), Inches(4.7), Inches(1.35), Inches(0.07), GOLD)


def build_approach(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_blank_bg(slide, OFF_WHITE)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), GOLD)
    add_logo(slide, Inches(0.5), Inches(0.28), Inches(2.2))
    add_textbox(
        slide, Inches(3.0), Inches(0.28), Inches(9.8), Inches(0.55),
        "Current Approach", 28, CHARCOAL, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.045), GOLD)
    add_footer(slide, 2)

    steps = [
        ("1", "Take the Excel file"),
        ("2", "Read the URLs from the Text column"),
        ("3", "Check whether each link opens"),
        ("4", "Mark it Working or Not Working"),
        ("5", "Save the result in a new column"),
    ]
    y = Inches(1.45)
    for num, text in steps:
        add_card(slide, Inches(1.6), y, Inches(10.1), Inches(0.92))
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.85), y + Inches(0.26), Inches(0.40), Inches(0.40)
        )
        fill_shape(badge, GOLD)
        add_textbox(
            slide, Inches(1.85), y + Inches(0.26), Inches(0.40), Inches(0.40),
            num, 16, CHARCOAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
        )
        add_textbox(
            slide, Inches(2.5), y, Inches(8.8), Inches(0.92),
            text, 22, CHARCOAL, False, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
        )
        y += Inches(1.05)


def build_future(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_blank_bg(slide, OFF_WHITE)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), GOLD)
    add_logo(slide, Inches(0.5), Inches(0.28), Inches(2.2))
    add_textbox(
        slide, Inches(3.0), Inches(0.28), Inches(9.8), Inches(0.55),
        "How We Can Expand This", 28, CHARCOAL, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
    )
    add_rect(slide, Inches(0.5), Inches(1.0), Inches(12.3), Inches(0.045), GOLD)
    add_footer(slide, 3)

    add_textbox(
        slide, Inches(0.7), Inches(1.25), Inches(12.0), Inches(0.55),
        "We will build a web page for this.",
        24, CHARCOAL, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
    )

    points = [
        ("Upload", "Put the Excel file on the web page."),
        ("Check", "Check the URLs from the Text column."),
        ("See result", "Show Working or Not Working on the screen."),
        ("Download", "Give back the Excel file with the result."),
    ]
    positions = [
        (Inches(0.7), Inches(2.05)),
        (Inches(6.85), Inches(2.05)),
        (Inches(0.7), Inches(4.45)),
        (Inches(6.85), Inches(4.45)),
    ]
    for (x, y), (title, body) in zip(positions, points):
        add_card(slide, x, y, Inches(5.75), Inches(2.05))
        add_rect(slide, x, y, Inches(0.12), Inches(2.05), GOLD)
        add_textbox(slide, x + Inches(0.4), y + Inches(0.22), Inches(5.1), Inches(0.45), title, 20, CHARCOAL, True)
        add_textbox(slide, x + Inches(0.4), y + Inches(0.75), Inches(5.1), Inches(0.95), body, 16, MUTED, False)


def main():
    make_logo()
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    build_title(prs)
    build_approach(prs)
    build_future(prs)
    prs.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
