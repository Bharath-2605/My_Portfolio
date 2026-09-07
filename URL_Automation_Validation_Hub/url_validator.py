#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import sys
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlparse

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

TEXT_COLUMN = 5
STATUS_COLUMN = 7
STATUS_HEADER = "URL Status"
URL_PATTERN = re.compile(
    r"(https?://[^\s<>\"'\)\]]+|www\.[^\s<>\"'\)\]]+)",
    re.IGNORECASE,
)
HYPERLINK_PATTERN = re.compile(
    r'HYPERLINK\s*\(\s*"([^"]+)"',
    re.IGNORECASE,
)
BROWSER_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/128.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Accept-Encoding": "gzip, deflate",
    "Connection": "keep-alive",
    "Upgrade-Insecure-Requests": "1",
}
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WORKING_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
WORKING_FONT = Font(color="006100")
FAILED_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
FAILED_FONT = Font(color="9C0006")
NEUTRAL_FILL = PatternFill(fill_type="solid", fgColor="FFEB9C")
NEUTRAL_FONT = Font(color="9C5700")
THREAD_LOCAL = threading.local()


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="url_validator",
        description="Validate URLs from column E of an Excel sheet and write results to column G.",
    )
    parser.add_argument("excel_file", help="Path to the .xlsx file")
    parser.add_argument(
        "-o",
        "--output",
        help="Output .xlsx path. Defaults to <original>_validated.xlsx",
    )
    parser.add_argument(
        "--sheet",
        help="Worksheet name or 1-based index. Defaults to the first sheet",
    )
    parser.add_argument(
        "--timeout",
        type=float,
        default=15.0,
        help="Request timeout in seconds (default: 15)",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=8,
        help="Number of concurrent URL checks (default: 8)",
    )
    return parser.parse_args()


def normalize_url(raw_url: str) -> str:
    url = raw_url.strip().rstrip(".,;:)]}'\"")
    if url.lower().startswith("www."):
        url = "https://" + url
    return url


def is_valid_http_url(url: str) -> bool:
    parsed = urlparse(url)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def extract_urls(cell_value: object, hyperlink_target: str | None) -> list[str]:
    collected: list[str] = []
    if hyperlink_target:
        collected.append(normalize_url(str(hyperlink_target)))
    if cell_value is None:
        text = ""
    else:
        text = str(cell_value).strip()
    if text.lower() in {"", "none", "nan"}:
        return unique_urls(collected)
    hyperlink_match = HYPERLINK_PATTERN.search(text)
    if hyperlink_match:
        collected.append(normalize_url(hyperlink_match.group(1)))
    found = URL_PATTERN.findall(text)
    if found:
        collected.extend(normalize_url(item) for item in found)
    elif " " not in text and "." in text and not text.startswith("="):
        collected.append(normalize_url(text))
    return unique_urls(collected)


def unique_urls(urls: list[str]) -> list[str]:
    seen: set[str] = set()
    unique: list[str] = []
    for url in urls:
        if url and url not in seen:
            seen.add(url)
            unique.append(url)
    return unique


def get_session() -> requests.Session:
    session = getattr(THREAD_LOCAL, "session", None)
    if session is None:
        session = requests.Session()
        session.headers.update(BROWSER_HEADERS)
        retry = Retry(
            total=2,
            connect=2,
            read=1,
            backoff_factor=0.4,
            status_forcelist=(429, 502, 503, 504),
            allowed_methods=frozenset({"GET", "HEAD"}),
            raise_on_status=False,
        )
        adapter = HTTPAdapter(max_retries=retry, pool_connections=16, pool_maxsize=16)
        session.mount("http://", adapter)
        session.mount("https://", adapter)
        THREAD_LOCAL.session = session
    return session


def classify_status_code(status_code: int) -> str:
    if 200 <= status_code < 400:
        return "working"
    return f"not working-{status_code} error"


def check_single_url(url: str, timeout: float) -> str:
    if not is_valid_http_url(url):
        return "not working-invalid URL"
    session = get_session()
    request_timeout = (min(10.0, timeout), timeout)
    try:
        with session.get(
            url,
            timeout=request_timeout,
            allow_redirects=True,
            stream=True,
        ) as response:
            return classify_status_code(response.status_code)
    except requests.exceptions.Timeout:
        return "not working-timeout error"
    except requests.exceptions.SSLError:
        return "not working-SSL error"
    except requests.exceptions.TooManyRedirects:
        return "not working-redirect error"
    except requests.exceptions.InvalidURL:
        return "not working-invalid URL"
    except requests.exceptions.MissingSchema:
        return "not working-invalid URL"
    except requests.exceptions.ConnectionError:
        return "not working-connection error"
    except requests.RequestException as error:
        reason = type(error).__name__.replace("Error", "").replace("Exception", "")
        reason = reason.lower() or "request"
        return f"not working-{reason} error"


def combine_statuses(results: list[str]) -> str:
    if not results:
        return "no URL found"
    unique = unique_urls(results)
    if len(unique) == 1:
        return unique[0]
    return " | ".join(unique)


def apply_status_style(cell, status: str) -> None:
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if status == "working" or status.startswith("working |"):
        cell.fill = WORKING_FILL
        cell.font = WORKING_FONT
    elif status == "no URL found":
        cell.fill = NEUTRAL_FILL
        cell.font = NEUTRAL_FONT
    else:
        cell.fill = FAILED_FILL
        cell.font = FAILED_FONT


def last_used_row(worksheet) -> int:
    last_row = 1
    for row in range(1, worksheet.max_row + 1):
        has_value = False
        for column in range(1, 7):
            value = worksheet.cell(row=row, column=column).value
            if value not in (None, ""):
                has_value = True
                break
        if has_value:
            last_row = row
    return last_row


def resolve_worksheet(workbook, sheet_argument: str | None):
    sheets = workbook.worksheets
    if not sheets:
        raise ValueError("The workbook does not contain any worksheets.")
    if sheet_argument is None:
        return sheets[0]
    if sheet_argument.isdigit():
        index = int(sheet_argument)
        if index < 1 or index > len(sheets):
            raise ValueError(
                f"Sheet index {index} is out of range. This file has {len(sheets)} sheet(s)."
            )
        return sheets[index - 1]
    if sheet_argument not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise ValueError(
            f'Sheet "{sheet_argument}" was not found. Available sheets: {available}'
        )
    return workbook[sheet_argument]


def collect_row_jobs(worksheet, last_row: int) -> list[tuple[int, list[str]]]:
    jobs: list[tuple[int, list[str]]] = []
    for row in range(2, last_row + 1):
        cell = worksheet.cell(row=row, column=TEXT_COLUMN)
        hyperlink_target = None
        if cell.hyperlink is not None and cell.hyperlink.target:
            hyperlink_target = str(cell.hyperlink.target)
        urls = extract_urls(cell.value, hyperlink_target)
        jobs.append((row, urls))
    return jobs


def validate_jobs(
    jobs: list[tuple[int, list[str]]],
    timeout: float,
    workers: int,
) -> dict[int, str]:
    statuses: dict[int, str] = {}
    pending: list[tuple[int, int, str]] = []
    row_slots: dict[int, list[str]] = {}
    for row, urls in jobs:
        if not urls:
            statuses[row] = "no URL found"
        else:
            row_slots[row] = [""] * len(urls)
            for index, url in enumerate(urls):
                pending.append((row, index, url))
    if not pending:
        return statuses
    worker_count = max(1, min(workers, len(pending)))
    completed = 0
    total = len(pending)
    print(f"Checking {total} URL(s) across {len(row_slots)} row(s)...")
    with ThreadPoolExecutor(max_workers=worker_count) as executor:
        future_map = {
            executor.submit(check_single_url, url, timeout): (row, index, url)
            for row, index, url in pending
        }
        for future in as_completed(future_map):
            row, index, url = future_map[future]
            try:
                result = future.result()
            except Exception as error:
                result = f"not working-{type(error).__name__} error"
            row_slots[row][index] = result
            completed += 1
            print(f"[{completed}/{total}] Row {row}: {url} -> {result}")
    for row, results in row_slots.items():
        statuses[row] = combine_statuses(results)
    return statuses


def write_statuses(worksheet, last_row: int, statuses: dict[int, str]) -> None:
    header_cell = worksheet.cell(row=1, column=STATUS_COLUMN)
    header_cell.value = STATUS_HEADER
    header_cell.fill = HEADER_FILL
    header_cell.font = HEADER_FONT
    header_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.column_dimensions["G"].width = 36
    for row in range(2, last_row + 1):
        status = statuses.get(row, "no URL found")
        cell = worksheet.cell(row=row, column=STATUS_COLUMN)
        cell.value = status
        apply_status_style(cell, status)


def default_output_path(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_validated{input_path.suffix}")


def load_excel(path: Path):
    try:
        return load_workbook(path)
    except FileNotFoundError:
        raise FileNotFoundError(f"Excel file not found: {path}") from None
    except PermissionError:
        raise PermissionError(
            f"Cannot read {path}. Close the file if it is open in Excel and try again."
        ) from None
    except InvalidFileException:
        raise ValueError(
            f"{path} is not a valid .xlsx file. Save the workbook as Excel Workbook (*.xlsx) and retry."
        ) from None
    except Exception as error:
        raise RuntimeError(f"Unable to open {path}: {error}") from error


def save_excel(workbook, path: Path) -> None:
    try:
        workbook.save(path)
    except PermissionError:
        raise PermissionError(
            f"Cannot save {path}. Close the file if it is open in Excel and try again."
        ) from None


def main() -> int:
    arguments = parse_arguments()
    input_path = Path(arguments.excel_file).expanduser().resolve()
    if input_path.suffix.lower() != ".xlsx":
        print("Please provide an .xlsx Excel file.", file=sys.stderr)
        return 1
    if not input_path.exists():
        print(f"Excel file not found: {input_path}", file=sys.stderr)
        return 1
    output_path = (
        Path(arguments.output).expanduser().resolve()
        if arguments.output
        else default_output_path(input_path)
    )
    if arguments.timeout <= 0:
        print("Timeout must be greater than 0.", file=sys.stderr)
        return 1
    if arguments.workers < 1:
        print("Workers must be at least 1.", file=sys.stderr)
        return 1
    try:
        workbook = load_excel(input_path)
        worksheet = resolve_worksheet(workbook, arguments.sheet)
        last_row = last_used_row(worksheet)
        if last_row < 2:
            print("No data rows were found under the header row.", file=sys.stderr)
            return 1
        jobs = collect_row_jobs(worksheet, last_row)
        statuses = validate_jobs(jobs, arguments.timeout, arguments.workers)
        write_statuses(worksheet, last_row, statuses)
        save_excel(workbook, output_path)
    except KeyboardInterrupt:
        print("\nValidation cancelled.", file=sys.stderr)
        return 130
    except Exception as error:
        print(str(error), file=sys.stderr)
        return 1
    working = sum(1 for status in statuses.values() if status == "working")
    missing = sum(1 for status in statuses.values() if status == "no URL found")
    failed = len(statuses) - working - missing
    print(f"Saved results to: {output_path}")
    print(f"Rows checked: {len(statuses)}")
    print(f"Working: {working}")
    print(f"Not working: {failed}")
    print(f"No URL found: {missing}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
